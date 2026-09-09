"""ParsingService:5 种文本格式 → 纯文本(禁 OCR,表格降级为文本流)。

契约(memory-bank/ai-design.md):
  输入 {file: bytes, filename: string} → 输出 {doc_id, title, text, warnings}
失败策略:解析失败(损坏/不支持)→ ParsingError 入日志不入索引;
          空文本 → EmptyTextError 拒绝并提示。
"""

import html as html_module
import io
import re
from dataclasses import dataclass
from pathlib import Path

import pymupdf as fitz  # PyMuPDF
from docx import Document as DocxDocument


class ParsingError(Exception):
    """解析失败(损坏/不支持格式)。"""


class EmptyTextError(ParsingError):
    """解析结果为空文本,拒绝入库。"""


@dataclass(frozen=True)
class ParseResult:
    doc_id: str
    title: str
    text: str
    warnings: tuple[str, ...] = ()


SUPPORTED_EXTENSIONS = (".md", ".txt", ".html", ".pdf", ".docx")

_HTML_COMMENT_RE = re.compile(r"<!--[\s\S]*?-->")
_HTML_TAG_RE = re.compile(r"<[^>]+>")
_HTML_TITLE_RE = re.compile(r"<title[^>]*>([^<]+)</title>", re.IGNORECASE)


def _clean_common(text: str) -> str:
    """通用清洗:行首尾去空白;压缩连续空行(>2 行 → 2 行)。"""
    lines = [line.strip() for line in text.splitlines()]
    out: list[str] = []
    blanks = 0
    for line in lines:
        if line:
            blanks = 0
            out.append(line)
        else:
            blanks += 1
            if blanks <= 2:
                out.append("")
    return "\n".join(out).strip()


def _extract_md_title(text: str) -> str:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("# "):
            return stripped[2:].strip()
    return ""


def _parse_md(content: bytes, filename: str) -> tuple[str, str, list[str]]:
    text = content.decode("utf-8", errors="replace")
    text = _HTML_COMMENT_RE.sub("", text)
    title = _extract_md_title(text)
    lines = [line for line in text.splitlines() if not line.strip().startswith("```")]
    return title, "\n".join(lines), []


def _parse_txt(content: bytes, filename: str) -> tuple[str, str, list[str]]:
    text = content.decode("utf-8", errors="replace")
    title = ""
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("synthetic") or stripped.startswith("<!--"):
            continue
        title = stripped
        break
    return title, text, []


def _parse_html(content: bytes, filename: str) -> tuple[str, str, list[str]]:
    text = content.decode("utf-8", errors="replace")
    title_match = _HTML_TITLE_RE.search(text)
    title = title_match.group(1).strip() if title_match else ""
    # 表格降级:tr → 行,td/th → 单元格以 " | " 分隔
    text = re.sub(r"(?i)</tr>", "\n", text)
    text = re.sub(r"(?i)</t[dh]>", " | ", text)
    # 块级元素结尾 → 换行
    text = re.sub(r"(?i)</(p|div|h\d|li|br)>", "\n", text)
    text = _HTML_COMMENT_RE.sub("", text)
    text = re.sub(r"(?i)<script[\s\S]*?</script>", "", text)
    text = re.sub(r"(?i)<style[\s\S]*?</style>", "", text)
    text = _HTML_TAG_RE.sub("", text)
    text = html_module.unescape(text)
    # 行级清理:去首尾空白,并去掉 </td> 替换遗留的行尾 " | "
    text = "\n".join(
        re.sub(r"\s*\|\s*$", "", line.strip()) for line in text.splitlines()
    )
    return title, text, []


def _parse_pdf(content: bytes, filename: str) -> tuple[str, str, list[str]]:
    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ParsingError(f"{filename}: PDF 打开失败: {exc}") from exc
    warnings: list[str] = []
    pages: list[str] = []
    try:
        if doc.page_count == 0:
            raise ParsingError(f"{filename}: PDF 无页面")
        for page in doc:
            pages.append(page.get_text("text"))
        if doc.needs_pass:
            warnings.append("PDF 加密,提取内容可能不完整")
    finally:
        doc.close()
    text = "\n".join(pages)
    title = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")
    return title, text, warnings


def _parse_docx(content: bytes, filename: str) -> tuple[str, str, list[str]]:
    try:
        doc = DocxDocument(io.BytesIO(content))
    except Exception as exc:
        raise ParsingError(f"{filename}: DOCX 打开失败: {exc}") from exc
    title = ""
    heading_title = ""
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        if not paragraph.text.strip():
            continue
        style_name = (paragraph.style.name or "") if paragraph.style else ""
        if not heading_title and "heading" in style_name.lower():
            heading_title = paragraph.text.strip()
        if not title:
            title = paragraph.text.strip()
        parts.append(paragraph.text)
    # 标题优先取 Heading 样式段落(真实文档标题常为 Heading 1)
    title = heading_title or title
    # 表格降级为文本流:每行 = 单元格 " | " 连接
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))
    return title, "\n".join(parts), []


_PARSERS = {
    ".md": _parse_md,
    ".txt": _parse_txt,
    ".html": _parse_html,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
}


def parse_file(filename: str, content: bytes) -> ParseResult:
    """统一入口:按扩展名分发,失败策略统一。"""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParsingError(
            f"不支持的格式: {ext}(支持: {', '.join(SUPPORTED_EXTENSIONS)})"
        )
    try:
        title, text, warnings = _PARSERS[ext](content, filename)
    except ParsingError:
        raise
    except Exception as exc:  # 解析器内部异常统一转为可判定失败
        raise ParsingError(f"{filename}: 解析异常: {exc}") from exc
    text = _clean_common(text)
    if not text.strip():
        raise EmptyTextError(f"{filename}: 解析结果为空文本,拒绝入库")
    return ParseResult(
        doc_id=Path(filename).stem,
        title=title or Path(filename).stem,
        text=text,
        warnings=tuple(warnings),
    )
